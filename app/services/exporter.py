import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

class BiExporter:
    @staticmethod
    def export_sales_to_excel(orders_data, start_date, end_date):
        """Gera planilha Excel (.xlsx) com o relatório de vendas."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Relatório de Vendas"

        # Cores padrão Bling
        bling_green = PatternFill(start_color="008B54", end_color="008B54", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="008B54")
        bold_font = Font(name="Calibri", size=10, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='DDDDDD'),
            right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'),
            bottom=Side(style='thin', color='DDDDDD')
        )

        # Cabeçalho do documento
        ws.append(["BLING BI - RELATÓRIO DE VENDAS E PEDIDOS"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Período: {start_date.strftime('%d/%m/%Y')} até {end_date.strftime('%d/%m/%Y')} | Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
        ws.append([])

        # Tabela
        headers = ["Nº Pedido", "Data", "Cliente", "Canal / Loja", "Itens", "Desconto (R$)", "Frete (R$)", "Total (R$)", "Situação"]
        ws.append(headers)

        header_row = 4
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.fill = bling_green
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        total_sum = 0.0
        row_idx = 5
        for o in orders_data:
            tot = float(o.get('total', 0.0))
            if o.get('status') == 'Atendido':
                total_sum += tot

            ws.append([
                o.get('number', ''),
                o.get('order_date', ''),
                o.get('customer_name', ''),
                o.get('channel', ''),
                o.get('items_count', 0),
                float(o.get('discount', 0.0)),
                float(o.get('shipping', 0.0)),
                tot,
                o.get('status', '')
            ])
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = thin_border
                if col_idx in [6, 7, 8]:
                    c.number_format = 'R$ #,##0.00'
                    c.alignment = Alignment(horizontal="right")
                elif col_idx in [1, 2, 5, 9]:
                    c.alignment = Alignment(horizontal="center")
            row_idx += 1

        # Linha Totalizadora
        ws.append(["TOTAL ATENDIDO", "", "", "", "", "", "", total_sum, ""])
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.font = bold_font
            c.border = thin_border
            if col_idx == 8:
                c.number_format = 'R$ #,##0.00'
                c.alignment = Alignment(horizontal="right")

        # Ajuste automático de largura de colunas
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_products_to_excel(products_data, start_date, end_date):
        """Gera planilha Excel com a Curva ABC e Produtos."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Curva ABC Produtos"

        bling_green = PatternFill(start_color="008B54", end_color="008B54", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="008B54")
        thin_border = Border(
            left=Side(style='thin', color='DDDDDD'),
            right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'),
            bottom=Side(style='thin', color='DDDDDD')
        )

        ws.append(["BLING BI - CURVA ABC DE PRODUTOS"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Período: {start_date.strftime('%d/%m/%Y')} até {end_date.strftime('%d/%m/%Y')}"])
        ws.append([])

        headers = ["Classe", "SKU", "Produto", "Qtd Vendida", "Faturamento (R$)", "Custo Total (R$)", "Margem Líquida (R$)", "% Margem", "% Acumulado", "Estoque Atual"]
        ws.append(headers)

        header_row = 4
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.fill = bling_green
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        row_idx = 5
        for p in products_data:
            ws.append([
                p.get('class', 'C'),
                p.get('sku', ''),
                p.get('name', ''),
                p.get('qty_sold', 0),
                p.get('revenue', 0.0),
                p.get('cost', 0.0),
                p.get('margin', 0.0),
                f"{p.get('margin_pct', 0.0):.1f}%",
                f"{p.get('accumulated_pct', 0.0):.1f}%",
                p.get('stock', 0)
            ])
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = thin_border
                if col_idx in [5, 6, 7]:
                    c.number_format = 'R$ #,##0.00'
                    c.alignment = Alignment(horizontal="right")
                elif col_idx in [1, 4, 8, 9, 10]:
                    c.alignment = Alignment(horizontal="center")
            row_idx += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_sales_to_pdf(orders_data, start_date, end_date, total_gross, avg_ticket):
        """Gera relatório de vendas profissional em PDF."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )

        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=18,
            textColor=colors.HexColor('#008B54'),
            spaceAfter=6
        )
        subtitle_style = ParagraphStyle(
            'SubTitleStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            textColor=colors.HexColor('#64748B'),
            spaceAfter=15
        )
        kpi_style = ParagraphStyle(
            'KpiStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=11,
            textColor=colors.HexColor('#1E293B')
        )

        elements.append(Paragraph("BlingBI - Relatório Executivo de Vendas", title_style))
        period_text = f"Período: {start_date.strftime('%d/%m/%Y')} a {end_date.strftime('%d/%m/%Y')} | Emitido em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        elements.append(Paragraph(period_text, subtitle_style))

        # Tabela de Resumo KPI
        kpi_data = [
            [
                Paragraph(f"<b>Faturamento Total:</b> R$ {total_gross:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'), kpi_style),
                Paragraph(f"<b>Pedidos Atendidos:</b> {len([o for o in orders_data if o.get('status') == 'Atendido'])}", kpi_style),
                Paragraph(f"<b>Ticket Médio:</b> R$ {avg_ticket:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'), kpi_style)
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[260, 260, 260])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
            ('PADDING', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('LINEBELOW', (0, 0), (-1, -1), 2, colors.HexColor('#008B54')),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 15))

        # Tabela Detalhada de Pedidos
        table_headers = ["Nº Pedido", "Data", "Cliente", "Canal / Loja", "Itens", "Total", "Status"]
        table_rows = [table_headers]

        for o in orders_data[:40]: # Limita a 40 para visualização limpa
            table_rows.append([
                str(o.get('number', '')),
                str(o.get('order_date', '')),
                str(o.get('customer_name', 'Consumidor'))[:25],
                str(o.get('channel', ''))[:18],
                str(o.get('items_count', 0)),
                f"R$ {float(o.get('total', 0.0)):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'),
                str(o.get('status', ''))
            ])

        orders_table = Table(table_rows, colWidths=[70, 75, 230, 160, 50, 100, 95])
        orders_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#008B54')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (0, 0), (1, -1), 'CENTER'),
            ('ALIGN', (4, 0), (4, -1), 'CENTER'),
            ('ALIGN', (5, 0), (5, -1), 'RIGHT'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        elements.append(orders_table)

        doc.build(elements)
        output.seek(0)
        return output

    @staticmethod
    def export_fashion_abc_to_excel(customers_data, start_date, end_date):
        """Exporta relatório da Curva ABC de Clientes com classificação 80/20 para Excel."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Curva ABC Clientes"

        bling_green = PatternFill(start_color="008B54", end_color="008B54", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="008B54")
        bold_font = Font(name="Calibri", size=10, bold=True)

        ws.append(["BLING BI - CURVA ABC DE CLIENTES (REGRA 80/20)"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Período: {start_date.strftime('%d/%m/%Y')} até {end_date.strftime('%d/%m/%Y')} | Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
        ws.append([])

        headers = ["Cliente", "Telefone", "Cidade / UF", "Pedidos", "Peças", "Receita (R$)", "% Receita", "% Acumulado", "Classe ABC"]
        ws.append(headers)

        header_row = 4
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.fill = bling_green
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        fill_a = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Verde claro
        fill_b = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid") # Amarelo claro
        fill_c = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid") # Cinza claro

        row_idx = 5
        for c in customers_data:
            abc = c.get('abc_class', 'C')
            ws.append([
                c.get('customer_name', ''),
                c.get('customer_phone', ''),
                f"{c.get('customer_city', '')}/{c.get('customer_state', '')}",
                c.get('orders_count', 0),
                c.get('total_pieces', 0),
                float(c.get('total_revenue', 0.0)),
                float(c.get('revenue_perc', 0.0)),
                float(c.get('cum_perc', 0.0)),
                abc
            ])
            cell_class = ws.cell(row=row_idx, column=9)
            cell_class.alignment = Alignment(horizontal="center")
            cell_class.font = bold_font
            if abc == 'A':
                cell_class.fill = fill_a
            elif abc == 'B':
                cell_class.fill = fill_b
            else:
                cell_class.fill = fill_c

            ws.cell(row=row_idx, column=6).number_format = 'R$ #,##0.00'
            ws.cell(row=row_idx, column=7).number_format = '0.0"%"'
            ws.cell(row=row_idx, column=8).number_format = '0.0"%"'
            row_idx += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_inactive_customers_to_excel(customers_data, reference_date):
        """Exporta lista de clientes inativos e recência para ação de Customer Success / Vendas."""
        wb = Workbook()
        ws = wb.active
        ws.title = "CRM Inativos & Recência"

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
        bold_font = Font(name="Calibri", size=10, bold=True)

        ws.append(["BLING BI - RELATÓRIO DE CLIENTES INATIVOS & RECÊNCIA (CRM)"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Data de Referência: {reference_date.strftime('%d/%m/%Y')} | Ação recomendada: Contato via WhatsApp"])
        ws.append([])

        headers = ["Cliente", "WhatsApp / Telefone", "Cidade / UF", "Tamanho Preferido", "Estilo", "Última Compra", "Dias Inativo", "Faixa / Risco", "LTV Histórico (R$)"]
        ws.append(headers)

        header_row = 4
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        row_idx = 5
        for c in customers_data:
            ws.append([
                c.get('customer_name', ''),
                c.get('customer_phone', ''),
                c.get('location', ''),
                c.get('preferred_size', 'M'),
                c.get('preferred_style', 'Casual'),
                c.get('last_purchase', ''),
                c.get('days_inactive', 0),
                c.get('tier', ''),
                float(c.get('lifetime_value', 0.0))
            ])
            ws.cell(row=row_idx, column=7).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=8).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=9).number_format = 'R$ #,##0.00'
            row_idx += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_conditional_bags_to_excel(bags_data, start_date, end_date):
        """Exporta relatório de Malas Condicionais e conversão de prova em casa."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Malas Condicionais"

        bling_green = PatternFill(start_color="008B54", end_color="008B54", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="008B54")

        ws.append(["BLING BI - RELATÓRIO DE MALAS CONDICIONAIS"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Período: {start_date.strftime('%d/%m/%Y')} até {end_date.strftime('%d/%m/%Y')}"])
        ws.append([])

        headers = ["Código Bag", "Data Envio", "Cliente", "Vendedora", "Peças Enviadas", "Peças Faturadas", "Peças Devolvidas", "Taxa Conversão", "Valor Faturado (R$)", "Situação"]
        ws.append(headers)

        header_row = 4
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.fill = bling_green
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        row_idx = 5
        for b in bags_data:
            ws.append([
                b.get('bag_code', ''),
                b.get('sent_date', ''),
                b.get('customer_name', ''),
                b.get('seller_name', ''),
                b.get('items_sent', 0),
                b.get('items_invoiced', 0),
                b.get('items_returned', 0),
                float(b.get('conversion_rate', 0.0)),
                float(b.get('total_invoiced_value', 0.0)),
                b.get('status', '')
            ])
            ws.cell(row=row_idx, column=8).number_format = '0.0"%"'
            ws.cell(row=row_idx, column=9).number_format = 'R$ #,##0.00'
            row_idx += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_fashion_abc_to_pdf(customers_data, summary, start_date, end_date):
        """Gera relatório da Curva ABC de Clientes em formato PDF profissional."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'FTitle', parent=styles['Normal'],
            fontName='Helvetica-Bold', fontSize=18,
            textColor=colors.HexColor('#008B54'), spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            'FSub', parent=styles['Normal'],
            fontName='Helvetica', fontSize=10,
            textColor=colors.HexColor('#64748B'), spaceAfter=14
        )
        kpi_style = ParagraphStyle(
            'FKpi', parent=styles['Normal'],
            fontName='Helvetica-Bold', fontSize=10,
            textColor=colors.HexColor('#1E293B')
        )

        elements.append(Paragraph("BlingBI - Relatório Estratégico: Curva ABC de Clientes (Regra 80/20)", title_style))
        elements.append(Paragraph(
            f"Período: {start_date.strftime('%d/%m/%Y')} a {end_date.strftime('%d/%m/%Y')} | Emitido em: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            subtitle_style
        ))

        # Tabela KPI de Resumo das Classes
        kpi_data = [
            [
                Paragraph(f"<b>Classe A:</b> {summary.get('class_a_count', 0)} clientes ({summary.get('class_a_perc', 0)}% da receita)", kpi_style),
                Paragraph(f"<b>Classe B:</b> {summary.get('class_b_count', 0)} clientes ({summary.get('class_b_perc', 0)}% da receita)", kpi_style),
                Paragraph(f"<b>Classe C:</b> {summary.get('class_c_count', 0)} clientes ({summary.get('class_c_perc', 0)}% da receita)", kpi_style),
                Paragraph(f"<b>Receita Total:</b> R$ {summary.get('total_revenue', 0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'), kpi_style)
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[195, 195, 195, 195])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
            ('PADDING', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('LINEBELOW', (0, 0), (-1, -1), 2, colors.HexColor('#008B54')),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 15))

        # Tabela de Clientes
        table_headers = ["Cliente", "Contato", "Cidade / UF", "Pedidos", "Peças", "Faturamento", "% Acumulado", "Classe"]
        table_rows = [table_headers]

        for c in customers_data[:45]:
            table_rows.append([
                str(c.get('customer_name', ''))[:24],
                str(c.get('customer_phone', '-')),
                f"{c.get('customer_city', '')}/{c.get('customer_state', '')}"[:18],
                str(c.get('orders_count', 0)),
                str(c.get('total_pieces', 0)),
                f"R$ {float(c.get('total_revenue', 0.0)):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'),
                f"{float(c.get('cum_perc', 0.0)):.1f}%",
                f"Classe {c.get('abc_class', 'C')}"
            ])

        client_table = Table(table_rows, colWidths=[170, 95, 115, 60, 60, 110, 85, 85])
        client_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#008B54')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8.5),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (3, 0), (4, -1), 'CENTER'),
            ('ALIGN', (5, 0), (6, -1), 'RIGHT'),
            ('ALIGN', (7, 0), (7, -1), 'CENTER'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        elements.append(client_table)

        doc.build(elements)
        output.seek(0)
        return output

    @staticmethod
    def export_conditional_bags_to_pdf(bags_data, start_date, end_date):
        """Gera relatório de Malas Condicionais em formato PDF profissional."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'BTitle', parent=styles['Normal'],
            fontName='Helvetica-Bold', fontSize=18,
            textColor=colors.HexColor('#008B54'), spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            'BSub', parent=styles['Normal'],
            fontName='Helvetica', fontSize=10,
            textColor=colors.HexColor('#64748B'), spaceAfter=14
        )
        kpi_style = ParagraphStyle(
            'BKpi', parent=styles['Normal'],
            fontName='Helvetica-Bold', fontSize=10,
            textColor=colors.HexColor('#1E293B')
        )

        elements.append(Paragraph("BlingBI - Relatório de Malas Condicionais & Prova em Casa", title_style))
        elements.append(Paragraph(
            f"Período: {start_date.strftime('%d/%m/%Y')} a {end_date.strftime('%d/%m/%Y')} | Emitido em: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            subtitle_style
        ))

        kpi_data = [
            [
                Paragraph(f"<b>Total Malas:</b> {bags_data.get('total_bags', 0)}", kpi_style),
                Paragraph(f"<b>Peças Enviadas:</b> {bags_data.get('total_sent', 0)}", kpi_style),
                Paragraph(f"<b>Peças Faturadas:</b> {bags_data.get('total_invoiced', 0)}", kpi_style),
                Paragraph(f"<b>Conversão da Bag:</b> {bags_data.get('overall_conversion_rate', 0)}%", kpi_style),
                Paragraph(f"<b>Receita:</b> R$ {bags_data.get('total_invoiced_value', 0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'), kpi_style)
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[156, 156, 156, 156, 156])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
            ('PADDING', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('LINEBELOW', (0, 0), (-1, -1), 2, colors.HexColor('#008B54')),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 15))

        table_headers = ["Código", "Data Envio", "Cliente", "Vendedora", "Enviadas", "Faturadas", "Devolvidas", "Conversão", "Faturado (R$)"]
        table_rows = [table_headers]

        for b in bags_data.get('bags', [])[:45]:
            table_rows.append([
                str(b.get('bag_code', '')),
                str(b.get('sent_date', '')),
                str(b.get('customer_name', ''))[:22],
                str(b.get('seller_name', ''))[:18],
                str(b.get('items_sent', 0)),
                str(b.get('items_invoiced', 0)),
                str(b.get('items_returned', 0)),
                f"{float(b.get('conversion_rate', 0.0)):.1f}%",
                f"R$ {float(b.get('total_invoiced_value', 0.0)):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
            ])

        bags_table = Table(table_rows, colWidths=[90, 75, 175, 130, 60, 60, 60, 65, 95])
        bags_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#008B54')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8.5),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (4, 0), (7, -1), 'CENTER'),
            ('ALIGN', (8, 0), (8, -1), 'RIGHT'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        elements.append(bags_table)

        doc.build(elements)
        output.seek(0)
        return output

    @staticmethod
    def export_service_performance_to_pdf(service_kpis, start_date, end_date):
        """Gera relatório de Desempenho do Atendimento (P.A., Ticket, Conversão) em PDF."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'STitle', parent=styles['Normal'],
            fontName='Helvetica-Bold', fontSize=18,
            textColor=colors.HexColor('#008B54'), spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            'SSub', parent=styles['Normal'],
            fontName='Helvetica', fontSize=10,
            textColor=colors.HexColor('#64748B'), spaceAfter=14
        )
        kpi_style = ParagraphStyle(
            'SKpi', parent=styles['Normal'],
            fontName='Helvetica-Bold', fontSize=10,
            textColor=colors.HexColor('#1E293B')
        )

        elements.append(Paragraph("BlingBI - Relatório de Desempenho do Atendimento & Vendas", title_style))
        elements.append(Paragraph(
            f"Período: {start_date.strftime('%d/%m/%Y')} a {end_date.strftime('%d/%m/%Y')} | Emitido em: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            subtitle_style
        ))

        kpi_data = [
            [
                Paragraph(f"<b>P.A. Geral da Loja:</b> {service_kpis.get('store_pa', 0):.2f} peças/venda", kpi_style),
                Paragraph(f"<b>Ticket Médio Geral:</b> R$ {service_kpis.get('store_ticket', 0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'), kpi_style),
                Paragraph(f"<b>Taxa de Conversão:</b> {service_kpis.get('store_conversion', 0):.1f}%", kpi_style)
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[260, 260, 260])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
            ('PADDING', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('LINEBELOW', (0, 0), (-1, -1), 2, colors.HexColor('#008B54')),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 15))

        table_headers = ["Consultora / Vendedora", "Atendimentos", "Vendas Fechadas", "P.A. (Peças)", "Ticket Médio", "Conversão", "Faturamento Total"]
        table_rows = [table_headers]

        for s in service_kpis.get('sellers_kpis', []):
            table_rows.append([
                str(s.get('seller_name', '')),
                str(s.get('total_attendances', 0)),
                str(s.get('orders_count', 0)),
                f"{float(s.get('pa', 0.0)):.2f}",
                f"R$ {float(s.get('ticket_medio', 0.0)):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'),
                f"{float(s.get('conversion_rate', 0.0)):.1f}%",
                f"R$ {float(s.get('total_revenue', 0.0)):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
            ])

        sellers_table = Table(table_rows, colWidths=[180, 95, 95, 95, 110, 95, 110])
        sellers_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#008B54')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (1, 0), (3, -1), 'CENTER'),
            ('ALIGN', (4, 0), (4, -1), 'RIGHT'),
            ('ALIGN', (5, 0), (5, -1), 'CENTER'),
            ('ALIGN', (6, 0), (6, -1), 'RIGHT'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        elements.append(sellers_table)

        doc.build(elements)
        output.seek(0)
        return output

    @staticmethod
    def export_inactives_to_pdf(inactives_data, reference_date):
        """Gera relatório de Clientes Inativos & Recência em formato PDF para ação no WhatsApp."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'ITitle', parent=styles['Normal'],
            fontName='Helvetica-Bold', fontSize=18,
            textColor=colors.HexColor('#1E293B'), spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            'ISub', parent=styles['Normal'],
            fontName='Helvetica', fontSize=10,
            textColor=colors.HexColor('#64748B'), spaceAfter=14
        )
        kpi_style = ParagraphStyle(
            'IKpi', parent=styles['Normal'],
            fontName='Helvetica-Bold', fontSize=10,
            textColor=colors.HexColor('#1E293B')
        )

        elements.append(Paragraph("BlingBI - Relatório de Recência & Clientes Inativos (CRM)", title_style))
        elements.append(Paragraph(
            f"Referência: {reference_date.strftime('%d/%m/%Y')} | Ação Estratégica: Contato consultivo via WhatsApp",
            subtitle_style
        ))

        kpi_data = [
            [
                Paragraph(f"<b>🟢 Ativos (<60d):</b> {inactives_data.get('active_count', 0)}", kpi_style),
                Paragraph(f"<b>🟡 Atenção (60-90d):</b> {inactives_data.get('warning_count', 0)}", kpi_style),
                Paragraph(f"<b>🟠 Inativos (91-180d):</b> {inactives_data.get('inactive_count', 0)}", kpi_style),
                Paragraph(f"<b>🔴 Churn (>180d):</b> {inactives_data.get('churn_count', 0)}", kpi_style)
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[195, 195, 195, 195])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
            ('PADDING', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('LINEBELOW', (0, 0), (-1, -1), 2, colors.HexColor('#1E293B')),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 15))

        table_headers = ["Cliente", "WhatsApp / Telefone", "Localidade", "Tamanho", "Estilo", "Última Compra", "Dias Inativo", "Faixa", "LTV Histórico"]
        table_rows = [table_headers]

        for c in inactives_data.get('customers', [])[:45]:
            table_rows.append([
                str(c.get('customer_name', ''))[:22],
                str(c.get('customer_phone', '-')),
                str(c.get('location', ''))[:16],
                str(c.get('preferred_size', 'M')),
                str(c.get('preferred_style', 'Casual'))[:14],
                str(c.get('last_purchase', '-')),
                f"{c.get('days_inactive', 0)} dias",
                str(c.get('tier', ''))[:14],
                f"R$ {float(c.get('lifetime_value', 0.0)):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
            ])

        inactives_table = Table(table_rows, colWidths=[140, 95, 95, 55, 80, 75, 75, 75, 90])
        inactives_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8.5),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (3, 0), (3, -1), 'CENTER'),
            ('ALIGN', (5, 0), (7, -1), 'CENTER'),
            ('ALIGN', (8, 0), (8, -1), 'RIGHT'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        elements.append(inactives_table)

        doc.build(elements)
        output.seek(0)
        return output

    # =========================================================================
    # EXPORTAÇÕES: RELATÓRIOS PADRONIZADOS ESTRATÉGICOS
    # =========================================================================

    @staticmethod
    def export_multichannel_to_excel(channels_data, summary, start_date, end_date):
        """Gera planilha Excel (.xlsx) com a Performance Multicanal & Margem por Canal."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Multicanal & Margem"

        emerald_green = PatternFill(start_color="008B54", end_color="008B54", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="008B54")
        bold_font = Font(name="Calibri", size=10, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='DDDDDD'),
            right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'),
            bottom=Side(style='thin', color='DDDDDD')
        )

        ws.append(["BLING BI - RELATÓRIO DE PERFORMANCE MULTICANAL & RENTABILIDADE"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Período: {start_date.strftime('%d/%m/%Y')} até {end_date.strftime('%d/%m/%Y')} | Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
        ws.append([])

        # Resumo Executivo
        ws.append(["Receita Líquida Total", "CMV Total", "Lucro Bruto Total", "Margem Geral (%)", "Pedidos Atendidos", "Fill Rate (%)"])
        ws.append([
            summary.get('total_net_sales', 0.0),
            summary.get('total_cmv', 0.0),
            summary.get('total_gross_profit', 0.0),
            f"{summary.get('overall_margin_pct', 0.0)}%",
            summary.get('orders_attended', 0),
            f"{summary.get('fill_rate_pct', 0.0)}%"
        ])
        for c in range(1, 7):
            cell = ws.cell(row=4, column=c)
            cell.font = header_font
            cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            val_cell = ws.cell(row=5, column=c)
            val_cell.font = bold_font
            if c in [1, 2, 3]:
                val_cell.number_format = 'R$ #,##0.00'
        ws.append([])

        # Tabela Detalhada por Canal
        headers = ["Canal / Loja", "Pedidos Atendidos", "Receita Líquida (R$)", "CMV (R$)", "Lucro Bruto (R$)", "Margem (%)", "Ticket Médio (R$)", "Peças / Pedido (P.A.)", "Cancelados", "Taxa Cancelamento (%)"]
        ws.append(headers)
        h_row = 7
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=h_row, column=col_idx)
            cell.fill = emerald_green
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center" if col_idx not in [1] else "left")

        current_row = h_row + 1
        for ch in channels_data:
            ws.append([
                ch.get('channel', 'Loja Virtual'),
                ch.get('orders_count', 0),
                ch.get('net_sales', 0.0),
                ch.get('cmv', 0.0),
                ch.get('gross_profit', 0.0),
                f"{ch.get('margin_pct', 0.0)}%",
                ch.get('avg_ticket', 0.0),
                ch.get('items_per_order', 0.0),
                ch.get('cancelled_count', 0),
                f"{ch.get('cancellation_rate', 0.0)}%"
            ])
            ws.cell(row=current_row, column=3).number_format = 'R$ #,##0.00'
            ws.cell(row=current_row, column=4).number_format = 'R$ #,##0.00'
            ws.cell(row=current_row, column=5).number_format = 'R$ #,##0.00'
            ws.cell(row=current_row, column=7).number_format = 'R$ #,##0.00'
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=current_row, column=col_idx).border = thin_border
            current_row += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_multichannel_to_pdf(channels_data, summary, start_date, end_date):
        """Gera PDF executivo da Performance Multicanal & Margem por Canal."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=25,
            leftMargin=25,
            topMargin=25,
            bottomMargin=25
        )
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#008B54'),
            spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            'DocSub',
            parent=styles['Normal'],
            fontSize=9.5,
            textColor=colors.HexColor('#64748B'),
            spaceAfter=12
        )
        kpi_style = ParagraphStyle('KPI', fontName='Helvetica', fontSize=8.5, textColor=colors.HexColor('#1E293B'), alignment=1)

        elements.append(Paragraph("<b>BlingBI | Relatório de Performance Multicanal & Rentabilidade</b>", title_style))
        elements.append(Paragraph(f"Período: {start_date.strftime('%d/%m/%Y')} até {end_date.strftime('%d/%m/%Y')} | Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style))

        def fmt_br(v):
            return f"R$ {float(v or 0.0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')

        kpi_data = [
            [
                Paragraph(f"<b>Receita Líquida:</b> {fmt_br(summary.get('total_net_sales', 0))}", kpi_style),
                Paragraph(f"<b>CMV Total:</b> {fmt_br(summary.get('total_cmv', 0))}", kpi_style),
                Paragraph(f"<b>Lucro Bruto:</b> {fmt_br(summary.get('total_gross_profit', 0))}", kpi_style),
                Paragraph(f"<b>Margem Bruta Geral:</b> {summary.get('overall_margin_pct', 0)}%", kpi_style),
                Paragraph(f"<b>Fill Rate:</b> {summary.get('fill_rate_pct', 0)}%", kpi_style)
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[156, 156, 156, 156, 156])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
            ('PADDING', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('LINEBELOW', (0, 0), (-1, -1), 2, colors.HexColor('#008B54')),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 14))

        table_headers = ["Canal / Marketplace", "Pedidos", "Receita Líquida", "CMV", "Lucro Bruto", "Margem %", "Ticket Médio", "P.A.", "Cancelados", "Taxa Canc."]
        table_rows = [table_headers]

        for ch in channels_data:
            table_rows.append([
                str(ch.get('channel', ''))[:22],
                str(ch.get('orders_count', 0)),
                fmt_br(ch.get('net_sales', 0)),
                fmt_br(ch.get('cmv', 0)),
                fmt_br(ch.get('gross_profit', 0)),
                f"{ch.get('margin_pct', 0)}%",
                fmt_br(ch.get('avg_ticket', 0)),
                str(ch.get('items_per_order', 0)),
                str(ch.get('cancelled_count', 0)),
                f"{ch.get('cancellation_rate', 0)}%"
            ])

        ch_table = Table(table_rows, colWidths=[130, 50, 85, 75, 85, 60, 80, 45, 65, 65])
        ch_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#008B54')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8.5),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (1, 0), (1, -1), 'CENTER'),
            ('ALIGN', (2, 0), (4, -1), 'RIGHT'),
            ('ALIGN', (5, 0), (5, -1), 'CENTER'),
            ('ALIGN', (6, 0), (6, -1), 'RIGHT'),
            ('ALIGN', (7, 0), (-1, -1), 'CENTER'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4.5),
            ('TOPPADDING', (0, 0), (-1, -1), 4.5),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        elements.append(ch_table)

        doc.build(elements)
        output.seek(0)
        return output

    @staticmethod
    def export_stock_coverage_to_excel(items_data, summary, start_date, end_date):
        """Gera planilha Excel (.xlsx) com a Curva ABC de Produtos, Cobertura em Dias e Ruptura."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Cobertura & Curva ABC"

        emerald_green = PatternFill(start_color="008B54", end_color="008B54", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="008B54")
        bold_font = Font(name="Calibri", size=10, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='DDDDDD'), right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'), bottom=Side(style='thin', color='DDDDDD')
        )

        ws.append(["BLING BI - RELATÓRIO DE CURVA ABC & DIAS DE COBERTURA DE ESTOQUE"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Período: {start_date.strftime('%d/%m/%Y')} até {end_date.strftime('%d/%m/%Y')} | Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
        ws.append([])

        ws.append(["Total de Itens Ativos", "Itens em Ruptura Crítica", "Capital Imobilizado (Deadstock)", "Capital Total em Estoque", "Cobertura Média da Loja"])
        ws.append([
            summary.get('total_products', 0),
            summary.get('critical_rupture_count', 0),
            summary.get('deadstock_capital', 0.0),
            summary.get('total_stock_capital', 0.0),
            f"{summary.get('avg_coverage_days', 0.0)} dias"
        ])
        for c in range(1, 6):
            cell = ws.cell(row=4, column=c)
            cell.font = header_font
            cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            val_cell = ws.cell(row=5, column=c)
            val_cell.font = bold_font
            if c in [3, 4]:
                val_cell.number_format = 'R$ #,##0.00'
        ws.append([])

        headers = ["SKU", "Produto", "Categoria", "Curva ABC", "Estoque Físico", "Custo Unit. (R$)", "Venda Média/Dia", "Cobertura (Dias)", "Status Estoque", "Receita Período (R$)", "Margem Bruta (R$)"]
        ws.append(headers)
        h_row = 7
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=h_row, column=col_idx)
            cell.fill = emerald_green
            cell.font = header_font

        current_row = h_row + 1
        for item in items_data:
            ws.append([
                item.get('sku', 'N/A'),
                item.get('name', ''),
                item.get('category', 'Geral'),
                item.get('abc_class', 'C'),
                item.get('stock', 0.0),
                item.get('cost', 0.0),
                item.get('run_rate', 0.0),
                item.get('coverage_days', 0.0),
                item.get('cov_status', 'EQUILIBRADO'),
                item.get('revenue', 0.0),
                item.get('gross_profit', 0.0)
            ])
            ws.cell(row=current_row, column=6).number_format = 'R$ #,##0.00'
            ws.cell(row=current_row, column=10).number_format = 'R$ #,##0.00'
            ws.cell(row=current_row, column=11).number_format = 'R$ #,##0.00'
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=current_row, column=col_idx).border = thin_border
            current_row += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_stock_coverage_to_pdf(items_data, summary, start_date, end_date):
        """Gera PDF executivo de Curva ABC e Cobertura de Estoque."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output, pagesize=landscape(A4),
            rightMargin=25, leftMargin=25, topMargin=25, bottomMargin=25
        )
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=16, textColor=colors.HexColor('#008B54'), spaceAfter=4)
        subtitle_style = ParagraphStyle('DocSub', parent=styles['Normal'], fontSize=9.5, textColor=colors.HexColor('#64748B'), spaceAfter=12)
        kpi_style = ParagraphStyle('KPI', fontName='Helvetica', fontSize=8.5, textColor=colors.HexColor('#1E293B'), alignment=1)

        elements.append(Paragraph("<b>BlingBI | Relatório de Curva ABC & Dias de Cobertura de Estoque</b>", title_style))
        elements.append(Paragraph(f"Período: {start_date.strftime('%d/%m/%Y')} até {end_date.strftime('%d/%m/%Y')} | Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style))

        def fmt_br(v):
            return f"R$ {float(v or 0.0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')

        kpi_data = [
            [
                Paragraph(f"<b>Itens Ativos:</b> {summary.get('total_products', 0)}", kpi_style),
                Paragraph(f"<b>🔴 Ruptura Crítica:</b> {summary.get('critical_rupture_count', 0)}", kpi_style),
                Paragraph(f"<b>Capital Deadstock:</b> {fmt_br(summary.get('deadstock_capital', 0))}", kpi_style),
                Paragraph(f"<b>Capital Estoque Total:</b> {fmt_br(summary.get('total_stock_capital', 0))}", kpi_style),
                Paragraph(f"<b>Cobertura Média:</b> {summary.get('avg_coverage_days', 0)} dias", kpi_style)
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[156, 156, 156, 156, 156])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
            ('PADDING', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('LINEBELOW', (0, 0), (-1, -1), 2, colors.HexColor('#008B54')),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 14))

        table_headers = ["SKU", "Produto", "Categoria", "Curva", "Estoque", "Custo Unit.", "Giro Diário", "Cobertura", "Status", "Receita Período"]
        table_rows = [table_headers]

        for it in items_data[:45]:
            table_rows.append([
                str(it.get('sku', 'N/A'))[:12],
                str(it.get('name', ''))[:22],
                str(it.get('category', 'Geral'))[:14],
                str(it.get('abc_class', 'C')),
                str(int(it.get('stock', 0))),
                fmt_br(it.get('cost', 0)),
                f"{it.get('run_rate', 0):.2f}",
                f"{it.get('coverage_days', 0)}d",
                str(it.get('cov_status', 'EQUILIBRADO')),
                fmt_br(it.get('revenue', 0))
            ])

        cov_table = Table(table_rows, colWidths=[70, 150, 85, 45, 50, 75, 60, 60, 85, 85])
        cov_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#008B54')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8.5),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (3, 0), (4, -1), 'CENTER'),
            ('ALIGN', (5, 0), (5, -1), 'RIGHT'),
            ('ALIGN', (6, 0), (8, -1), 'CENTER'),
            ('ALIGN', (9, 0), (9, -1), 'RIGHT'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        elements.append(cov_table)

        doc.build(elements)
        output.seek(0)
        return output

    @staticmethod
    def export_dre_to_excel(dre_lines, summary, start_date, end_date):
        """Gera planilha Excel (.xlsx) da DRE Gerencial Sintética."""
        wb = Workbook()
        ws = wb.active
        ws.title = "DRE Gerencial"

        title_font = Font(name="Calibri", size=14, bold=True, color="008B54")
        bold_font = Font(name="Calibri", size=11, bold=True)
        regular_font = Font(name="Calibri", size=10)
        subtotal_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        header_fill = PatternFill(start_color="008B54", end_color="008B54", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        ws.append(["BLING BI - DEMONSTRATIVO DE RESULTADO DO EXERCÍCIO (DRE GERENCIAL)"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Período: {start_date.strftime('%d/%m/%Y')} até {end_date.strftime('%d/%m/%Y')} | Regime: Competência"])
        ws.append([])

        headers = ["Código", "Estrutura Contábil / Gerencial", "Valor (R$)", "Margem (%)"]
        ws.append(headers)
        for col_idx in range(1, 5):
            cell = ws.cell(row=4, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font

        current_row = 5
        for line in dre_lines:
            ws.append([
                line.get('code', ''),
                line.get('name', ''),
                line.get('value', 0.0),
                f"{line.get('margin_pct', '')}%" if 'margin_pct' in line else ''
            ])
            val_cell = ws.cell(row=current_row, column=3)
            val_cell.number_format = 'R$ #,##0.00'

            if line.get('type') in ['subtotal', 'total', 'positive']:
                ws.cell(row=current_row, column=1).font = bold_font
                ws.cell(row=current_row, column=2).font = bold_font
                val_cell.font = bold_font
                for c in range(1, 5):
                    ws.cell(row=current_row, column=c).fill = subtotal_fill

            current_row += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_dre_to_pdf(dre_lines, summary, start_date, end_date):
        """Gera PDF executivo formatado da DRE Gerencial Sintética."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output, pagesize=A4,
            rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30
        )
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=15, textColor=colors.HexColor('#008B54'), spaceAfter=4)
        subtitle_style = ParagraphStyle('DocSub', parent=styles['Normal'], fontSize=9.5, textColor=colors.HexColor('#64748B'), spaceAfter=14)

        elements.append(Paragraph("<b>BlingBI | Demonstrativo de Resultado do Exercício (DRE)</b>", title_style))
        elements.append(Paragraph(f"Período: {start_date.strftime('%d/%m/%Y')} até {end_date.strftime('%d/%m/%Y')} | Regime de Competência", subtitle_style))

        def fmt_br(v):
            return f"R$ {float(v or 0.0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')

        table_rows = [["Código", "Estrutura Gerencial / DRE", "Valor (R$)", "Margem %"]]

        for l in dre_lines:
            table_rows.append([
                str(l.get('code', '')),
                str(l.get('name', '')),
                fmt_br(l.get('value', 0.0)),
                f"{l.get('margin_pct', '')}%" if 'margin_pct' in l else '-'
            ])

        dre_table = Table(table_rows, colWidths=[55, 310, 110, 60])
        style_cmds = [
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('ALIGN', (0, 0), (0, -1), 'CENTER'),
            ('ALIGN', (1, 0), (1, -1), 'LEFT'),
            ('ALIGN', (2, 0), (2, -1), 'RIGHT'),
            ('ALIGN', (3, 0), (3, -1), 'CENTER'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]

        # Estilização condicional de linhas da DRE
        for idx, line in enumerate(dre_lines, start=1):
            if line.get('type') in ['subtotal', 'total']:
                style_cmds.append(('BACKGROUND', (0, idx), (-1, idx), colors.HexColor('#F1F5F9')))
                style_cmds.append(('FONTNAME', (0, idx), (-1, idx), 'Helvetica-Bold'))
            elif line.get('type') == 'positive':
                style_cmds.append(('FONTNAME', (0, idx), (-1, idx), 'Helvetica-Bold'))

        dre_table.setStyle(TableStyle(style_cmds))
        elements.append(dre_table)

        doc.build(elements)
        output.seek(0)
        return output

    @staticmethod
    def export_rfm_to_excel(customers_data, summary):
        """Gera planilha Excel (.xlsx) da Matriz RFM de Clientes com clusters acionáveis."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Matriz RFM Clientes"

        title_font = Font(name="Calibri", size=14, bold=True, color="008B54")
        bold_font = Font(name="Calibri", size=11, bold=True)
        header_fill = PatternFill(start_color="008B54", end_color="008B54", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='DDDDDD'), right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'), bottom=Side(style='thin', color='DDDDDD')
        )

        ws.append(["BLING BI - MATRIZ RFM (RECÊNCIA, FREQUÊNCIA E VALOR) & SEGMENTAÇÃO CRM"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Data de Referência: {datetime.now().strftime('%d/%m/%Y')} | Total de Clientes Analisados: {summary.get('total_customers', 0)}"])
        ws.append([])

        ws.append(["Total de Clientes", "👑 Campeões", "💎 Leais", "⚠️ Em Risco", "💤 Churn / Inativos", "LTV Médio"])
        ws.append([
            summary.get('total_customers', 0),
            summary.get('champions_count', 0),
            summary.get('loyal_count', 0),
            summary.get('at_risk_count', 0),
            summary.get('churn_count', 0),
            summary.get('avg_ltv', 0.0)
        ])
        for c in range(1, 7):
            ws.cell(row=4, column=c).fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            ws.cell(row=4, column=c).font = header_font
            val_cell = ws.cell(row=5, column=c)
            val_cell.font = bold_font
            if c == 6:
                val_cell.number_format = 'R$ #,##0.00'
        ws.append([])

        headers = ["Cliente", "Telefone / WhatsApp", "E-mail", "Cidade / UF", "Última Compra", "Recência (Dias)", "Frequência", "Monetário (R$)", "Score RFM", "Segmento / Cluster", "Ação Sugerida"]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=7, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font

        current_row = 8
        for c in customers_data:
            ws.append([
                c.get('name', ''),
                c.get('phone', ''),
                c.get('email', ''),
                f"{c.get('city', '')} - {c.get('state', '')}",
                c.get('last_order_date', ''),
                c.get('recency_days', 0),
                c.get('frequency', 0),
                c.get('monetary', 0.0),
                c.get('rfm_score', ''),
                c.get('cluster', ''),
                c.get('action_tip', '')
            ])
            ws.cell(row=current_row, column=8).number_format = 'R$ #,##0.00'
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=current_row, column=col_idx).border = thin_border
            current_row += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_rfm_to_pdf(customers_data, summary):
        """Gera PDF executivo da Matriz RFM de Clientes."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output, pagesize=landscape(A4),
            rightMargin=25, leftMargin=25, topMargin=25, bottomMargin=25
        )
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=16, textColor=colors.HexColor('#008B54'), spaceAfter=4)
        subtitle_style = ParagraphStyle('DocSub', parent=styles['Normal'], fontSize=9.5, textColor=colors.HexColor('#64748B'), spaceAfter=12)
        kpi_style = ParagraphStyle('KPI', fontName='Helvetica', fontSize=8.5, textColor=colors.HexColor('#1E293B'), alignment=1)

        elements.append(Paragraph("<b>BlingBI | Matriz RFM & Segmentação Estratégica de Clientes</b>", title_style))
        elements.append(Paragraph(f"Base Analisada em: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Total de Clientes: {summary.get('total_customers', 0)}", subtitle_style))

        def fmt_br(v):
            return f"R$ {float(v or 0.0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')

        kpi_data = [
            [
                Paragraph(f"<b>👑 Campeões:</b> {summary.get('champions_count', 0)}", kpi_style),
                Paragraph(f"<b>💎 Leais:</b> {summary.get('loyal_count', 0)}", kpi_style),
                Paragraph(f"<b>⚠️ Em Risco:</b> {summary.get('at_risk_count', 0)}", kpi_style),
                Paragraph(f"<b>💤 Churn:</b> {summary.get('churn_count', 0)}", kpi_style),
                Paragraph(f"<b>LTV Médio:</b> {fmt_br(summary.get('avg_ltv', 0))}", kpi_style)
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[156, 156, 156, 156, 156])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
            ('PADDING', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('LINEBELOW', (0, 0), (-1, -1), 2, colors.HexColor('#008B54')),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 14))

        table_headers = ["Cliente", "Telefone", "Localidade", "Última Compra", "Recência", "Freq.", "LTV (R$)", "Score", "Segmento / Cluster"]
        table_rows = [table_headers]

        for c in customers_data[:45]:
            table_rows.append([
                str(c.get('name', ''))[:22],
                str(c.get('phone', '-')),
                f"{c.get('city', '')}/{c.get('state', '')}"[:16],
                str(c.get('last_order_date', '-')),
                f"{c.get('recency_days', 0)}d",
                str(c.get('frequency', 0)),
                fmt_br(c.get('monetary', 0)),
                str(c.get('rfm_score', '')),
                str(c.get('cluster', ''))[:20]
            ])

        rfm_table = Table(table_rows, colWidths=[130, 95, 100, 75, 55, 45, 90, 50, 140])
        rfm_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#008B54')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8.5),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (3, 0), (5, -1), 'CENTER'),
            ('ALIGN', (6, 0), (6, -1), 'RIGHT'),
            ('ALIGN', (7, 0), (7, -1), 'CENTER'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        elements.append(rfm_table)

        doc.build(elements)
        output.seek(0)
        return output

    @staticmethod
    def export_product_search_to_excel(products, summary, search_filters=None):
        """Exporta resultados da busca geral de produtos por tamanho, cor e nome para Excel."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Busca de Produtos"

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="008B54", end_color="008B54", fill_type="solid")
        title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
        kpi_label_font = Font(name="Calibri", size=9, color="64748B", bold=True)
        kpi_val_font = Font(name="Calibri", size=12, bold=True, color="008B54")
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws['A1'] = "BlingBI - Busca Geral de Produtos (Catálogo, Tamanho & Cor)"
        ws['A1'].font = title_font
        ws['A2'] = f"Exportado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws['A2'].font = Font(name="Calibri", size=10, color="64748B")

        # KPIs no topo
        ws['A4'] = "PRODUTOS ENCONTRADOS"
        ws['A4'].font = kpi_label_font
        ws['A5'] = summary.get('total_found', len(products))
        ws['A5'].font = kpi_val_font

        ws['C4'] = "PEÇAS EM ESTOQUE"
        ws['C4'].font = kpi_label_font
        ws['C5'] = summary.get('total_stock_qty', 0)
        ws['C5'].font = kpi_val_font

        ws['E4'] = "VALOR TOTAL ESTOQUE"
        ws['E4'].font = kpi_label_font
        ws['E5'] = f"R$ {summary.get('total_stock_value', 0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
        ws['E5'].font = kpi_val_font

        ws['G4'] = "PREÇO MÉDIO"
        ws['G4'].font = kpi_label_font
        ws['G5'] = f"R$ {summary.get('avg_price', 0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
        ws['G5'].font = kpi_val_font

        headers = ["Código / SKU", "Nome do Produto", "Tamanho", "Cor", "Estoque", "Preço Venda (R$)", "Custo (R$)", "Margem (R$)", "Margem (%)", "Status"]
        start_row = 7

        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center" if col_idx in [1, 3, 4, 5, 10] else ("right" if col_idx >= 6 else "left"), vertical="center")

        for r_idx, p in enumerate(products, start_row + 1):
            ws.cell(row=r_idx, column=1, value=p.get('sku', '')).alignment = Alignment(horizontal="center")
            ws.cell(row=r_idx, column=2, value=p.get('name', ''))
            ws.cell(row=r_idx, column=3, value=p.get('size', '')).alignment = Alignment(horizontal="center")
            ws.cell(row=r_idx, column=4, value=p.get('color', '')).alignment = Alignment(horizontal="center")
            
            stock_cell = ws.cell(row=r_idx, column=5, value=p.get('stock', 0))
            stock_cell.alignment = Alignment(horizontal="center")
            stock_cell.number_format = '#,##0'

            price_cell = ws.cell(row=r_idx, column=6, value=p.get('price', 0.0))
            price_cell.alignment = Alignment(horizontal="right")
            price_cell.number_format = 'R$ #,##0.00'

            cost_cell = ws.cell(row=r_idx, column=7, value=p.get('cost', 0.0))
            cost_cell.alignment = Alignment(horizontal="right")
            cost_cell.number_format = 'R$ #,##0.00'

            margin_cell = ws.cell(row=r_idx, column=8, value=p.get('margin_rs', 0.0))
            margin_cell.alignment = Alignment(horizontal="right")
            margin_cell.number_format = 'R$ #,##0.00'

            margin_pct_cell = ws.cell(row=r_idx, column=9, value=(p.get('margin_pct', 0.0) / 100.0))
            margin_pct_cell.alignment = Alignment(horizontal="right")
            margin_pct_cell.number_format = '0.0%'

            status_text = 'Ativo' if p.get('status') == 'A' else 'Inativo'
            ws.cell(row=r_idx, column=10, value=status_text).alignment = Alignment(horizontal="center")

            for col in range(1, 11):
                ws.cell(row=r_idx, column=col).border = thin_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_product_search_to_pdf(products, summary, search_filters=None):
        """Exporta resultados da busca geral de produtos por tamanho, cor e nome para PDF."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=14 * mm, leftMargin=14 * mm,
            topMargin=12 * mm, bottomMargin=12 * mm
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'PSTitle', parent=styles['Heading1'],
            fontName='Helvetica-Bold', fontSize=15,
            textColor=colors.HexColor('#008B54'), spaceAfter=4
        )
        sub_style = ParagraphStyle(
            'PSSub', parent=styles['Normal'],
            fontName='Helvetica', fontSize=9.5,
            textColor=colors.HexColor('#64748B'), spaceAfter=12
        )
        kpi_style = ParagraphStyle(
            'PSKpi', parent=styles['Normal'],
            fontName='Helvetica-Bold', fontSize=9.5,
            textColor=colors.HexColor('#1E293B')
        )

        elements = [
            Paragraph("BlingBI - Relatório de Busca e Catálogo de Produtos", title_style),
            Paragraph(f"Emitido em: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Produtos Catalogados", sub_style)
        ]

        # Resumo executivo
        def fmt_br(v):
            return f"R$ {v:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')

        kpi_row = [
            Paragraph(f"<b>Encontrados:</b> {summary.get('total_found', len(products))}", kpi_style),
            Paragraph(f"<b>Peças Estoque:</b> {int(summary.get('total_stock_qty', 0)):,}".replace(',', '.'), kpi_style),
            Paragraph(f"<b>Valor Estoque:</b> {fmt_br(summary.get('total_stock_value', 0))}", kpi_style),
            Paragraph(f"<b>Preço Médio:</b> {fmt_br(summary.get('avg_price', 0))}", kpi_style)
        ]
        kpi_tab = Table([kpi_row], colWidths=[195, 195, 195, 195])
        kpi_tab.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('PADDING', (0, 0), (-1, -1), 7),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('LINEBELOW', (0, 0), (-1, -1), 2, colors.HexColor('#008B54')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        elements.append(kpi_tab)
        elements.append(Spacer(1, 12))

        headers = ["SKU", "Nome do Produto", "Tam", "Cor", "Estoque", "Preço (R$)", "Custo (R$)", "Margem", "Status"]
        table_rows = [headers]

        for p in products[:150]:
            table_rows.append([
                str(p.get('sku', ''))[:14],
                str(p.get('name', ''))[:38],
                str(p.get('size', ''))[:6],
                str(p.get('color', ''))[:16],
                str(int(p.get('stock', 0))),
                fmt_br(p.get('price', 0)),
                fmt_br(p.get('cost', 0)),
                f"{p.get('margin_pct', 0.0):.1f}%",
                'Ativo' if p.get('status') == 'A' else 'Inativo'
            ])

        prod_table = Table(table_rows, colWidths=[70, 240, 45, 100, 55, 80, 80, 55, 55])
        prod_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#008B54')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (2, 0), (4, -1), 'CENTER'),
            ('ALIGN', (5, 0), (7, -1), 'RIGHT'),
            ('ALIGN', (8, 0), (8, -1), 'CENTER'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
            ('TOPPADDING', (0, 0), (-1, -1), 3.5),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        elements.append(prod_table)

        doc.build(elements)
        output.seek(0)
        return output

    @staticmethod
    def export_filtered_customers_to_excel(customers, summary, filter_meta=None):
        """Exporta lista de clientes filtrados por atributos e períodos para planilha Excel (.xlsx)."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Clientes CRM Filtrados"

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="008B54", end_color="008B54", fill_type="solid")
        title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
        kpi_label_font = Font(name="Calibri", size=9, color="64748B", bold=True)
        kpi_val_font = Font(name="Calibri", size=12, bold=True, color="008B54")
        bold_font = Font(name="Calibri", size=10, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        f_label = filter_meta.get('filter_label', 'Clientes') if filter_meta else 'Clientes'
        p_label = filter_meta.get('period_label', 'Período') if filter_meta else 'Período'

        ws['A1'] = f"BlingBI - Inteligência de Clientes e CRM: {f_label}"
        ws['A1'].font = title_font
        ws['A2'] = f"Período: {p_label} | Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws['A2'].font = Font(name="Calibri", size=10, color="64748B")

        # Linha de KPIs
        ws['A4'] = "CLIENTES FILTRADOS"
        ws['A4'].font = kpi_label_font
        ws['A5'] = summary.get('total_customers', len(customers))
        ws['A5'].font = kpi_val_font

        ws['C4'] = "FATURAMENTO DO GRUPO"
        ws['C4'].font = kpi_label_font
        ws['C5'] = f"R$ {summary.get('total_revenue', 0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
        ws['C5'].font = kpi_val_font

        ws['E4'] = "TOTAL DE PEÇAS"
        ws['E4'].font = kpi_label_font
        ws['E5'] = summary.get('total_items', 0)
        ws['E5'].font = kpi_val_font

        ws['G4'] = "TICKET MÉDIO"
        ws['G4'].font = kpi_label_font
        ws['G5'] = f"R$ {summary.get('avg_ticket', 0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
        ws['G5'].font = kpi_val_font

        ws['I4'] = "TAXA DE RECOMPRA"
        ws['I4'].font = kpi_label_font
        ws['I5'] = f"{summary.get('recurrence_rate', 0):.1f}%"
        ws['I5'].font = kpi_val_font

        headers = [
            "Cliente", "Tipo", "CPF / CNPJ", "Telefone", "Cidade / UF",
            "Pedidos (Período)", "Peças Compradas", "Total Gasto (R$)", "Ticket Médio (R$)",
            "Última Compra", "Dias sem Comprar", "Status Inatividade",
            "Bags Consignado", "Total Consignado (R$)", "Gasto Histórico Lifetime (R$)"
        ]

        header_row = 7
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        row_idx = 8
        for c in customers:
            ws.append([
                c.get('name', ''),
                c.get('person_type', 'F'),
                c.get('document', ''),
                c.get('phone', ''),
                f"{c.get('city', '')} / {c.get('state', '')}" if c.get('city') else c.get('state', ''),
                c.get('period_orders', 0),
                c.get('period_items', 0),
                float(c.get('period_spent', 0.0)),
                float(c.get('period_avg_ticket', 0.0)),
                c.get('last_order_date', '-'),
                c.get('days_since_last') if c.get('days_since_last') is not None else 'Sem compras',
                c.get('inactivity_label', '-'),
                c.get('bags_count', 0),
                float(c.get('bags_invoiced_value', 0.0)),
                float(c.get('lifetime_spent', 0.0))
            ])

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if col_idx in [8, 9, 14, 15]:
                    cell.number_format = 'R$ #,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx in [2, 6, 7, 10, 11, 12, 13]:
                    cell.alignment = Alignment(horizontal="center")

            row_idx += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_filtered_customers_to_pdf(customers, summary, filter_meta=None):
        """Exporta lista de clientes filtrados por atributos e períodos para PDF Executivo (Landscape A4)."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=10 * mm,
            leftMargin=10 * mm,
            topMargin=10 * mm,
            bottomMargin=10 * mm
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustTitle',
            parent=styles['Heading1'],
            fontSize=15,
            leading=18,
            textColor=colors.HexColor('#1E293B'),
            fontName='Helvetica-Bold'
        )
        subtitle_style = ParagraphStyle(
            'CustSub',
            parent=styles['Normal'],
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor('#64748B')
        )
        kpi_label_style = ParagraphStyle('CustKpiL', fontSize=7.5, leading=9, textColor=colors.HexColor('#64748B'), alignment=1)
        kpi_val_style = ParagraphStyle('CustKpiV', fontSize=10, leading=12, textColor=colors.HexColor('#008B54'), fontName='Helvetica-Bold', alignment=1)

        f_label = filter_meta.get('filter_label', 'Clientes') if filter_meta else 'Clientes'
        p_label = filter_meta.get('period_label', 'Período') if filter_meta else 'Período'

        elements = []
        elements.append(Paragraph(f"BlingBI — Inteligência de Clientes & CRM: {f_label}", title_style))
        elements.append(Paragraph(f"Período: <b>{p_label}</b> | Emissão: {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style))
        elements.append(Spacer(1, 4 * mm))

        # Tabela de KPIs
        fmt_br = lambda v: f"R$ {v:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
        kpi_data = [
            [
                Paragraph("CLIENTES FILTRADOS", kpi_label_style),
                Paragraph("FATURAMENTO TOTAL", kpi_label_style),
                Paragraph("TOTAL DE PEÇAS", kpi_label_style),
                Paragraph("TICKET MÉDIO", kpi_label_style),
                Paragraph("TAXA RECOMPRA", kpi_label_style)
            ],
            [
                Paragraph(str(summary.get('total_customers', len(customers))), kpi_val_style),
                Paragraph(fmt_br(summary.get('total_revenue', 0)), kpi_val_style),
                Paragraph(str(summary.get('total_items', 0)), kpi_val_style),
                Paragraph(fmt_br(summary.get('avg_ticket', 0)), kpi_val_style),
                Paragraph(f"{summary.get('recurrence_rate', 0):.1f}%", kpi_val_style)
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[150, 160, 150, 150, 150])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 4 * mm))

        headers = ["Cliente", "CPF / CNPJ", "Cidade / UF", "Ped.", "Peças", "Total Gasto", "Última Compra", "Dias Inat.", "Consignado"]
        table_rows = [headers]

        for c in customers[:70]:
            name_val = c.get('name', '')
            if len(name_val) > 28:
                name_val = name_val[:26] + '...'
            days_str = str(c.get('days_since_last')) if c.get('days_since_last') is not None else '-'
            consign_str = f"{c.get('bags_count', 0)} bags" if c.get('has_consignment') else '-'

            table_rows.append([
                name_val,
                c.get('document', '') or '-',
                f"{c.get('city', '')}/{c.get('state', '')}" if c.get('city') else c.get('state', '') or '-',
                str(c.get('period_orders', 0)),
                str(c.get('period_items', 0)),
                fmt_br(c.get('period_spent', 0)),
                c.get('last_order_date', '-'),
                days_str,
                consign_str
            ])

        col_widths = [180, 95, 100, 40, 45, 90, 75, 65, 70]
        cust_table = Table(table_rows, colWidths=col_widths)
        cust_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#008B54')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 7.5),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (3, 0), (4, -1), 'CENTER'),
            ('ALIGN', (5, 0), (5, -1), 'RIGHT'),
            ('ALIGN', (6, 0), (8, -1), 'CENTER'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        elements.append(cust_table)

        doc.build(elements)
        output.seek(0)
        return output

    @staticmethod
    def export_cash_statement_to_excel(movements, summary, daily_closings, start_date, end_date):
        """Exporta o Fechamento de Caixa e Extrato de Movimentação Financeira para planilha Excel (.xlsx)."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Extrato de Movimentação"

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="008B54", end_color="008B54", fill_type="solid")
        title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
        kpi_label_font = Font(name="Calibri", size=9, color="64748B", bold=True)
        kpi_val_font = Font(name="Calibri", size=12, bold=True, color="008B54")
        bold_font = Font(name="Calibri", size=10, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws['A1'] = "BlingBI - Fechamento de Caixa & Extrato de Movimentação Financeira"
        ws['A1'].font = title_font
        ws['A2'] = f"Período: {start_date.strftime('%d/%m/%Y')} a {end_date.strftime('%d/%m/%Y')} | Emitido em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws['A2'].font = Font(name="Calibri", size=10, color="64748B")

        # KPIs no Topo
        ws['A4'] = "SALDO INICIAL (ABERTURA)"
        ws['A4'].font = kpi_label_font
        ws['A5'] = f"R$ {summary.get('opening_balance', 0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
        ws['A5'].font = kpi_val_font

        ws['C4'] = "TOTAL ENTRADAS (+)"
        ws['C4'].font = kpi_label_font
        ws['C5'] = f"R$ {summary.get('total_inflows', 0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
        ws['C5'].font = Font(name="Calibri", size=12, bold=True, color="166534")

        ws['E4'] = "TOTAL SAÍDAS (-)"
        ws['E4'].font = kpi_label_font
        ws['E5'] = f"R$ {summary.get('total_outflows', 0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
        ws['E5'].font = Font(name="Calibri", size=12, bold=True, color="991B1B")

        ws['G4'] = "RESULTADO DO PERÍODO"
        ws['G4'].font = kpi_label_font
        ws['G5'] = f"R$ {summary.get('net_result', 0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
        ws['G5'].font = kpi_val_font

        ws['I4'] = "SALDO FINAL (FECHAMENTO)"
        ws['I4'].font = kpi_label_font
        ws['I5'] = f"R$ {summary.get('closing_balance', 0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
        ws['I5'].font = Font(name="Calibri", size=12, bold=True, color="1E293B")

        # Tabela de Movimentações
        headers = [
            "Data", "Tipo", "Documento", "Descrição da Operação",
            "Cliente / Fornecedor", "Categoria", "Forma de Pagamento",
            "Valor da Operação (R$)", "Saldo Acumulado (R$)", "Situação"
        ]

        header_row = 7
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        row_idx = 8
        for m in movements:
            val = float(m.get('amount', 0.0))
            is_in = (m.get('type') == 'IN')
            signed_val = val if is_in else -val

            ws.append([
                m.get('date_str', ''),
                'Entrada (+)' if is_in else 'Saída (-)',
                m.get('doc_number', ''),
                m.get('description', ''),
                m.get('entity', ''),
                m.get('category', ''),
                m.get('payment_method', ''),
                signed_val,
                float(m.get('running_balance', 0.0)),
                m.get('status', 'Liquidado')
            ])

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if col_idx in [8, 9]:
                    cell.number_format = 'R$ #,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                    if col_idx == 8:
                        cell.font = Font(name="Calibri", size=10, bold=True, color="166534" if is_in else "991B1B")
                elif col_idx in [1, 2, 3, 7, 10]:
                    cell.alignment = Alignment(horizontal="center")

            row_idx += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # Aba 2: Fechamentos Diários Consolidados
        ws_daily = wb.create_sheet(title="Fechamentos Diários")
        ws_daily['A1'] = "BlingBI - Histórico de Fechamentos de Caixa Diários"
        ws_daily['A1'].font = title_font
        ws_daily['A2'] = f"Período: {start_date.strftime('%d/%m/%Y')} a {end_date.strftime('%d/%m/%Y')}"
        ws_daily['A2'].font = Font(name="Calibri", size=10, color="64748B")

        daily_headers = ["Data", "Saldo Abertura (R$)", "Entradas do Dia (R$)", "Saídas do Dia (R$)", "Resultado Líquido (R$)", "Saldo Fechamento (R$)", "Operações", "Status"]
        for col_idx, h in enumerate(daily_headers, 1):
            cell = ws_daily.cell(row=4, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        d_row = 5
        for d in daily_closings:
            ws_daily.append([
                d.get('date_str', ''),
                float(d.get('opening_balance', 0.0)),
                float(d.get('total_in', 0.0)),
                float(d.get('total_out', 0.0)),
                float(d.get('net_day', 0.0)),
                float(d.get('closing_balance', 0.0)),
                d.get('operations_count', 0),
                d.get('status', 'Conciliado')
            ])
            for col_idx in range(1, len(daily_headers) + 1):
                cell = ws_daily.cell(row=d_row, column=col_idx)
                cell.border = thin_border
                if col_idx in [2, 3, 4, 5, 6]:
                    cell.number_format = 'R$ #,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx in [1, 7, 8]:
                    cell.alignment = Alignment(horizontal="center")
            d_row += 1

        for col in ws_daily.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws_daily.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_cash_statement_to_pdf(movements, summary, daily_closings, start_date, end_date):
        """Exporta o Fechamento de Caixa e Extrato de Movimentação Financeira para PDF Executivo (Landscape A4)."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=10 * mm,
            leftMargin=10 * mm,
            topMargin=10 * mm,
            bottomMargin=10 * mm
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CashTitle',
            parent=styles['Heading1'],
            fontSize=15,
            leading=18,
            textColor=colors.HexColor('#1E293B'),
            fontName='Helvetica-Bold'
        )
        subtitle_style = ParagraphStyle(
            'CashSub',
            parent=styles['Normal'],
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor('#64748B')
        )
        kpi_label_style = ParagraphStyle('CashKpiL', fontSize=7.5, leading=9, textColor=colors.HexColor('#64748B'), alignment=1)
        kpi_val_style = ParagraphStyle('CashKpiV', fontSize=10, leading=12, textColor=colors.HexColor('#008B54'), fontName='Helvetica-Bold', alignment=1)
        kpi_neg_style = ParagraphStyle('CashKpiN', fontSize=10, leading=12, textColor=colors.HexColor('#991B1B'), fontName='Helvetica-Bold', alignment=1)

        fmt_br = lambda v: f"R$ {v:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')

        elements = []
        elements.append(Paragraph("BlingBI — Fechamento de Caixa & Extrato de Movimentação Financeira", title_style))
        elements.append(Paragraph(f"Período: <b>{start_date.strftime('%d/%m/%Y')} até {end_date.strftime('%d/%m/%Y')}</b> | Emissão: {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style))
        elements.append(Spacer(1, 4 * mm))

        # Tabela de KPIs do Caixa
        kpi_data = [
            [
                Paragraph("SALDO INICIAL", kpi_label_style),
                Paragraph("TOTAL ENTRADAS (+)", kpi_label_style),
                Paragraph("TOTAL SAÍDAS (-)", kpi_label_style),
                Paragraph("RESULTADO LÍQUIDO", kpi_label_style),
                Paragraph("SALDO FECHAMENTO", kpi_label_style)
            ],
            [
                Paragraph(fmt_br(summary.get('opening_balance', 0)), kpi_val_style),
                Paragraph(fmt_br(summary.get('total_inflows', 0)), kpi_val_style),
                Paragraph(fmt_br(summary.get('total_outflows', 0)), kpi_neg_style),
                Paragraph(fmt_br(summary.get('net_result', 0)), kpi_val_style if summary.get('net_result', 0) >= 0 else kpi_neg_style),
                Paragraph(fmt_br(summary.get('closing_balance', 0)), kpi_val_style)
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[150, 160, 150, 150, 150])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 4 * mm))

        headers = ["Data", "Tipo", "Doc.", "Descrição da Operação", "Cliente / Fornecedor", "Meio Pagto.", "Valor (R$)", "Saldo Acumulado"]
        table_rows = [headers]

        for m in movements[:80]:
            is_in = (m.get('type') == 'IN')
            desc_val = m.get('description', '')
            if len(desc_val) > 28:
                desc_val = desc_val[:26] + '...'
            ent_val = m.get('entity', '')
            if len(ent_val) > 22:
                ent_val = ent_val[:20] + '...'

            val_str = f"+ {fmt_br(m.get('amount', 0))}" if is_in else f"- {fmt_br(m.get('amount', 0))}"

            table_rows.append([
                m.get('date_str', ''),
                'Entrada (+)' if is_in else 'Saída (-)',
                m.get('doc_number', ''),
                desc_val,
                ent_val,
                m.get('payment_method', ''),
                val_str,
                fmt_br(m.get('running_balance', 0))
            ])

        col_widths = [65, 65, 65, 190, 140, 95, 75, 75]
        statement_table = Table(table_rows, colWidths=col_widths)
        statement_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#008B54')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 7.5),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (0, 0), (2, -1), 'CENTER'),
            ('ALIGN', (5, 0), (5, -1), 'CENTER'),
            ('ALIGN', (6, 0), (7, -1), 'RIGHT'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        elements.append(statement_table)

        doc.build(elements)
        output.seek(0)
        return output

    @staticmethod
    def export_accounts_to_excel(entries, summary, filters=None):
        """Gera planilha Excel (.xlsx) profissional de Contas a Pagar e Contas a Receber."""
        wb = Workbook()
        
        # Aba 1: Lançamentos
        ws = wb.active
        ws.title = "Lançamentos"

        filters = filters or {}
        bling_green = PatternFill(start_color="008B54", end_color="008B54", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="008B54")
        bold_font = Font(name="Calibri", size=10, bold=True)
        gray_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        pay_font = Font(name="Calibri", size=10, color="991B1B", bold=True)
        rec_font = Font(name="Calibri", size=10, color="008B54", bold=True)

        thin_border = Border(
            left=Side(style='thin', color='DDDDDD'),
            right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'),
            bottom=Side(style='thin', color='DDDDDD')
        )

        ws.append(["BLING BI - RELATÓRIO DE CONTAS A PAGAR & RECEBER"])
        ws.cell(row=1, column=1).font = title_font

        tipo_filtro = filters.get('entry_type', 'TODOS')
        status_filtro = filters.get('status', 'TODOS')
        ini = filters.get('start_date', 'Início')
        fim = filters.get('end_date', 'Fim')
        ws.append([f"Período: {ini} até {fim} | Tipo: {tipo_filtro} | Situação: {status_filtro} | Emissão: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
        ws.append([])

        # KPIs no topo
        ws.append([
            "A Receber (Previsto)", summary.get('total_receber_pending', 0.0),
            "A Pagar (Previsto)", summary.get('total_pagar_pending', 0.0),
            "Saldo Previsto", summary.get('saldo_pending', 0.0),
            "Em Atraso", summary.get('total_overdue', 0.0)
        ])
        kpi_row = 4
        for c_idx in [1, 3, 5, 7]:
            ws.cell(row=kpi_row, column=c_idx).font = bold_font
            ws.cell(row=kpi_row, column=c_idx).fill = gray_fill
        for c_idx in [2, 4, 6, 8]:
            c = ws.cell(row=kpi_row, column=c_idx)
            c.font = bold_font
            c.number_format = 'R$ #,##0.00'
        ws.append([])

        headers = [
            "Tipo", "Vencimento", "Descrição", "Favorecido / Pagador",
            "Categoria", "Forma Pagto.", "Valor Nominal (R$)", "Desconto (R$)",
            "Juros/Multa (R$)", "Valor Final (R$)", "Situação", "Data Pagto.", "ID / Bling", "Observações"
        ]
        ws.append(headers)
        header_row = 6
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.fill = bling_green
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        row_idx = 7
        for e in entries:
            if hasattr(e, 'to_dict'):
                e = e.to_dict()
            elif not isinstance(e, dict):
                e = {
                    'entry_type': getattr(e, 'entry_type', ''),
                    'due_date_formatted': getattr(e, 'due_date_formatted', '') or str(getattr(e, 'due_date', '') or ''),
                    'description': getattr(e, 'description', ''),
                    'payee_payer': getattr(e, 'payee_payer', '') or getattr(e, 'entity_name', ''),
                    'category': getattr(e, 'category', ''),
                    'payment_method': getattr(e, 'payment_method', ''),
                    'amount': getattr(e, 'amount', 0.0),
                    'discount': getattr(e, 'discount', 0.0),
                    'fine_fee': getattr(e, 'fine_fee', 0.0),
                    'net_amount': getattr(e, 'net_amount', 0.0),
                    'status': getattr(e, 'status', ''),
                    'payment_date_formatted': getattr(e, 'payment_date_formatted', '') or str(getattr(e, 'payment_date', '') or ''),
                    'bling_id': getattr(e, 'bling_id', None) or getattr(e, 'id', None),
                    'notes': getattr(e, 'notes', '')
                }
            is_pagar = (e.get('entry_type') == 'PAGAR')
            val_nom = float(e.get('amount') or 0.0)
            desc = float(e.get('discount') or 0.0)
            fine = float(e.get('fine_fee') or 0.0)
            final_val = float(e.get('net_amount') or (val_nom - desc + fine))

            ws.append([
                'PAGAR' if is_pagar else 'RECEBER',
                e.get('due_date_formatted') or e.get('due_date') or '',
                e.get('description') or '',
                e.get('payee_payer') or '',
                e.get('category') or '',
                e.get('payment_method') or '',
                val_nom,
                desc,
                fine,
                final_val,
                e.get('status') or '',
                e.get('payment_date_formatted') or e.get('payment_date') or '',
                str(e.get('bling_id') or e.get('id') or ''),
                e.get('notes') or ''
            ])

            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = thin_border
                if col_idx == 1:
                    c.alignment = Alignment(horizontal="center")
                    c.font = pay_font if is_pagar else rec_font
                elif col_idx in [2, 11, 12, 13]:
                    c.alignment = Alignment(horizontal="center")
                elif col_idx in [7, 8, 9, 10]:
                    c.number_format = 'R$ #,##0.00'
                    c.alignment = Alignment(horizontal="right")
            row_idx += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

        # Aba 2: Resumo Financeiro
        ws2 = wb.create_sheet(title="Resumo & Alertas")
        ws2.append(["RESUMO GERAL DE CONTAS A PAGAR E RECEBER"])
        ws2.cell(row=1, column=1).font = title_font
        ws2.append([])

        resumo_items = [
            ("Total a Receber (Previsto / Pendente)", summary.get('total_receber_pending', 0.0)),
            ("Total a Receber (Liquidado / Pago)", summary.get('total_receber_paid', 0.0)),
            ("Total Geral a Receber", summary.get('total_receber', 0.0)),
            ("Total a Pagar (Previsto / Pendente)", summary.get('total_pagar_pending', 0.0)),
            ("Total a Pagar (Liquidado / Pago)", summary.get('total_pagar_paid', 0.0)),
            ("Total Geral a Pagar", summary.get('total_pagar', 0.0)),
            ("Saldo Líquido Previsto (Pendente)", summary.get('saldo_pending', 0.0)),
            ("Saldo Líquido Realizado (Pago)", summary.get('saldo_paid', 0.0)),
            ("Saldo Líquido Total", summary.get('saldo_total', 0.0)),
            ("Total de Títulos em Atraso (R$)", summary.get('total_overdue', 0.0)),
            ("Quantidade de Títulos em Atraso", summary.get('count_overdue', 0)),
            ("Títulos com Vencimento Hoje (Qtd)", summary.get('count_today', 0)),
            ("Títulos com Vencimento Hoje (R$)", summary.get('total_today', 0.0)),
            ("Títulos com Vencimento nos Próx. 7 Dias (Qtd)", summary.get('count_week', 0)),
            ("Títulos com Vencimento nos Próx. 7 Dias (R$)", summary.get('total_week', 0.0)),
            ("Títulos com Vencimento no Mês (Qtd)", summary.get('count_month', 0)),
            ("Títulos com Vencimento no Mês (R$)", summary.get('total_month', 0.0)),
        ]

        ws2.append(["Indicador", "Valor"])
        ws2.cell(row=3, column=1).fill = bling_green
        ws2.cell(row=3, column=1).font = header_font
        ws2.cell(row=3, column=2).fill = bling_green
        ws2.cell(row=3, column=2).font = header_font

        r_idx = 4
        for label, val in resumo_items:
            ws2.append([label, val])
            c1 = ws2.cell(row=r_idx, column=1)
            c2 = ws2.cell(row=r_idx, column=2)
            c1.border = thin_border
            c2.border = thin_border
            if isinstance(val, float):
                c2.number_format = 'R$ #,##0.00'
                c2.alignment = Alignment(horizontal="right")
            else:
                c2.alignment = Alignment(horizontal="center")
            r_idx += 1

        ws2.column_dimensions['A'].width = 45
        ws2.column_dimensions['B'].width = 22

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_accounts_to_pdf(entries, summary, filters=None):
        """Exporta relatório executivo de Contas a Pagar & Receber em PDF (Landscape A4)."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=10 * mm,
            leftMargin=10 * mm,
            topMargin=10 * mm,
            bottomMargin=10 * mm
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'AccTitle',
            parent=styles['Heading1'],
            fontSize=15,
            leading=18,
            textColor=colors.HexColor('#1E293B'),
            fontName='Helvetica-Bold'
        )
        subtitle_style = ParagraphStyle(
            'AccSub',
            parent=styles['Normal'],
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor('#64748B')
        )
        kpi_label_style = ParagraphStyle('AccKpiL', fontSize=7.5, leading=9, textColor=colors.HexColor('#64748B'), alignment=1)
        kpi_val_green = ParagraphStyle('AccKpiG', fontSize=10, leading=12, textColor=colors.HexColor('#008B54'), fontName='Helvetica-Bold', alignment=1)
        kpi_val_red = ParagraphStyle('AccKpiR', fontSize=10, leading=12, textColor=colors.HexColor('#991B1B'), fontName='Helvetica-Bold', alignment=1)
        kpi_val_blue = ParagraphStyle('AccKpiB', fontSize=10, leading=12, textColor=colors.HexColor('#2563EB'), fontName='Helvetica-Bold', alignment=1)

        fmt_br = lambda v: f"R$ {float(v or 0):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')

        filters = filters or {}
        tipo = filters.get('entry_type', 'TODOS')
        status = filters.get('status', 'TODOS')
        ini = filters.get('start_date', 'Início')
        fim = filters.get('end_date', 'Fim')

        elements = []
        elements.append(Paragraph("BlingBI — Relatório Executivo de Contas a Pagar & Receber", title_style))
        elements.append(Paragraph(f"Período: <b>{ini} até {fim}</b> | Tipo: <b>{tipo}</b> | Situação: <b>{status}</b> | Emissão: {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style))
        elements.append(Spacer(1, 4 * mm))

        # KPIs do Topo
        kpi_data = [
            [
                Paragraph("A RECEBER (PREVISTO)", kpi_label_style),
                Paragraph("A PAGAR (PREVISTO)", kpi_label_style),
                Paragraph("SALDO PREVISTO", kpi_label_style),
                Paragraph("SALDO REALIZADO", kpi_label_style),
                Paragraph("EM ATRASO", kpi_label_style)
            ],
            [
                Paragraph(fmt_br(summary.get('total_receber_pending', 0)), kpi_val_green),
                Paragraph(fmt_br(summary.get('total_pagar_pending', 0)), kpi_val_red),
                Paragraph(fmt_br(summary.get('saldo_pending', 0)), kpi_val_blue if summary.get('saldo_pending', 0) >= 0 else kpi_val_red),
                Paragraph(fmt_br(summary.get('saldo_paid', 0)), kpi_val_green if summary.get('saldo_paid', 0) >= 0 else kpi_val_red),
                Paragraph(fmt_br(summary.get('total_overdue', 0)), kpi_val_red)
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[150, 150, 160, 150, 160])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 4 * mm))

        # Tabela de Títulos
        headers = ["Tipo", "Vencimento", "Descrição / Histórico", "Favorecido / Pagador", "Categoria", "Valor Líquido (R$)", "Situação", "Pagamento"]
        table_rows = [headers]

        for e in entries[:150]:
            if hasattr(e, 'to_dict'):
                e = e.to_dict()
            elif not isinstance(e, dict):
                e = {
                    'entry_type': getattr(e, 'entry_type', ''),
                    'due_date_formatted': getattr(e, 'due_date_formatted', '') or str(getattr(e, 'due_date', '') or ''),
                    'description': getattr(e, 'description', ''),
                    'payee_payer': getattr(e, 'payee_payer', '') or getattr(e, 'entity_name', ''),
                    'category': getattr(e, 'category', ''),
                    'payment_method': getattr(e, 'payment_method', ''),
                    'amount': getattr(e, 'amount', 0.0),
                    'discount': getattr(e, 'discount', 0.0),
                    'fine_fee': getattr(e, 'fine_fee', 0.0),
                    'net_amount': getattr(e, 'net_amount', 0.0),
                    'status': getattr(e, 'status', ''),
                    'payment_date_formatted': getattr(e, 'payment_date_formatted', '') or str(getattr(e, 'payment_date', '') or ''),
                    'bling_id': getattr(e, 'bling_id', None) or getattr(e, 'id', None),
                    'notes': getattr(e, 'notes', '')
                }
            is_pagar = (e.get('entry_type') == 'PAGAR')
            val_nom = float(e.get('amount') or 0.0)
            desc = float(e.get('discount') or 0.0)
            fine = float(e.get('fine_fee') or 0.0)
            final_val = float(e.get('net_amount') or (val_nom - desc + fine))

            desc_txt = e.get('description') or ''
            if len(desc_txt) > 30:
                desc_txt = desc_txt[:28] + '...'
            ent_txt = e.get('payee_payer') or ''
            if len(ent_txt) > 24:
                ent_txt = ent_txt[:22] + '...'
            cat_txt = e.get('category') or ''
            if len(cat_txt) > 18:
                cat_txt = cat_txt[:16] + '...'

            table_rows.append([
                'PAGAR' if is_pagar else 'RECEBER',
                e.get('due_date_formatted') or e.get('due_date') or '',
                desc_txt,
                ent_txt,
                cat_txt,
                fmt_br(final_val),
                e.get('status') or '',
                e.get('payment_date_formatted') or e.get('payment_date') or '-'
            ])

        col_widths = [60, 65, 185, 150, 105, 85, 65, 65]
        acc_table = Table(table_rows, colWidths=col_widths)
        acc_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#008B54')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 7.5),
            ('ALIGN', (0, 0), (1, -1), 'CENTER'),
            ('ALIGN', (2, 0), (4, -1), 'LEFT'),
            ('ALIGN', (5, 0), (5, -1), 'RIGHT'),
            ('ALIGN', (6, 0), (7, -1), 'CENTER'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        elements.append(acc_table)

        doc.build(elements)
        output.seek(0)
        return output




